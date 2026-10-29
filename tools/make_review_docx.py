#!/usr/bin/env python3
"""把 申论/本次批改.md 排成可打印的 A4 讲义（docx + pdf）。

用法：
    python tools/make_review_docx.py            # 用默认源文件
    python tools/make_review_docx.py --src 申论/本次批改.md --out docs/print
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent))
from make_handout import (  # noqa: E402
    FONT,
    OUT,
    ROOT,
    fill_table,
    heading,
    para,
    set_style,
    to_pdf,
)
from docx.oxml.ns import qn  # noqa: E402

SRC = ROOT / "申论" / "本次批改.md"

MARK_RE = re.compile(r"==(?:重点|背景|意义|问题|做法|结论|释义)\|([^=\n]*?)==")
MARK_RE2 = re.compile(r"==([^=\n]*?)==")
LINK_RE = re.compile(r"\[([^\]]+)\]\([^)]*\)")


def plain(s: str) -> str:
    """去掉语义色、链接、行内代码、粗体标记，得到纯文本。"""
    s = MARK_RE.sub(r"\1", s)
    s = MARK_RE2.sub(r"\1", s)
    s = LINK_RE.sub(r"\1", s)
    s = s.replace("**", "").replace("`", "")
    return s.strip()


def add_rich(p, text: str, *, size: float = 10.5, color=None):
    """把 **粗体** / ==语义色== 拆成多段 run：语义色=加粗，粗体=加粗。"""
    text = LINK_RE.sub(r"\1", text)
    parts = re.split(r"(\*\*.+?\*\*|==[^=\n]*?==)", text)
    for part in parts:
        if not part:
            continue
        bold = False
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            part = part[2:-2]
            bold = True
        elif part.startswith("==") and part.endswith("==") and len(part) > 4:
            part = part[2:-2]
            bold = True
            if "|" in part:
                head, rest = part.split("|", 1)
                if head in ("重点", "背景", "意义", "问题", "做法", "结论", "释义"):
                    part = rest
        part = part.replace("`", "")
        run = p.add_run(part)
        run.bold = bold
        run.font.size = Pt(size)
        run.font.name = FONT
        run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        if color:
            run.font.color.rgb = RGBColor(*color)


def is_table_row(line: str) -> bool:
    return line.startswith("|") and line.count("|") >= 2


def split_row(line: str) -> list[str]:
    s = line.strip()
    # 语义色标记里可能带竖线（==重点|xxx==），先换成占位符，切完再还原
    s = re.sub(r"==(?:重点|背景|意义|问题|做法|结论|释义)\|[^=\n]*==",
               lambda m: m.group(0).replace("|", "\x01"), s)
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.replace("\x01", "|").strip() for c in s.split("|")]


def parse(src: Path) -> dict:
    lines = src.read_text(encoding="utf-8").splitlines()
    title = "本次批改"
    block: list[dict] = []
    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        s = line.strip()
        if not s:
            i += 1
            continue
        if s == "---":
            block.append({"t": "hr"})
            i += 1
            continue
        if s.startswith("# ") and title == "本次批改":
            title = plain(s[2:])
            i += 1
            continue
        if s.startswith("### "):
            block.append({"t": "h3", "text": plain(s[4:])})
            i += 1
            continue
        if s.startswith("## "):
            block.append({"t": "h2", "text": plain(s[3:])})
            i += 1
            continue
        if is_table_row(s):
            rows = []
            while i < len(lines) and is_table_row(lines[i].strip()):
                rows.append(split_row(lines[i]))
                i += 1
            rows = [r for r in rows if not all(re.fullmatch(r":?-{2,}:?", c or "-") for c in r)]
            if rows:
                block.append({"t": "table", "rows": rows})
            continue
        if s.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip()[1:].strip())
                i += 1
            block.append({"t": "quote", "text": " ".join(x for x in buf if x)})
            continue
        m = re.match(r"^(?:[-*])\s+(.*)$", s)
        if m:
            block.append({"t": "li", "text": m.group(1)})
            i += 1
            continue
        m = re.match(r"^(\d+)[.、]\s*(.*)$", s)
        if m:
            block.append({"t": "li", "num": m.group(1), "text": m.group(2)})
            i += 1
            continue
        block.append({"t": "p", "text": s})
        i += 1
    return {"title": title, "block": block}


def table_widths(rows: list[list[str]], total: float = 17.6) -> list[float]:
    n = max(len(r) for r in rows)
    if n <= 2:
        return [total * 0.42, total * 0.58][:n] if n == 2 else [total]
    head = total * 0.26
    rest = (total - head) / (n - 1)
    return [head] + [rest] * (n - 1)


def emit_blocks(doc, blocks) -> None:
    for b in blocks:
        t = b["t"]
        if t == "hr":
            continue
        if t == "h2":
            heading(doc, b["text"], size=12.5)
        elif t == "h3":
            para(doc, b["text"], size=11, bold=True, space_before=8, space_after=2)
        elif t == "quote":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.3
            add_rich(p, b["text"], size=10, color=(0x33, 0x33, 0x33))
            pPr = p._p.get_or_add_pPr()
            from docx.oxml import OxmlElement

            bdr = OxmlElement("w:pBdr")
            left = OxmlElement("w:left")
            left.set(qn("w:val"), "single")
            left.set(qn("w:sz"), "12")
            left.set(qn("w:space"), "6")
            left.set(qn("w:color"), "999999")
            bdr.append(left)
            pPr.append(bdr)
        elif t == "li":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.first_line_indent = Cm(-0.5)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.3
            prefix = f"{b['num']}. " if b.get("num") else "· "
            add_rich(p, prefix + b["text"], size=10.5)
        elif t == "table":
            rows = b["rows"]
            header = rows[0]
            n = max(len(r) for r in rows)
            header = header + [""] * (n - len(header))
            body = [[(r + [""] * n)[:n] for r in rows[1:]]]
            body = body[0] if body else []
            table = doc.add_table(rows=1, cols=n)
            fill_table(table, [plain(h) for h in header], [[plain(c) for c in row] for row in body],
                       table_widths(rows), size=9.5, header_size=9.5)
            para(doc, "", size=4, space_after=0)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.32
            add_rich(p, b["text"], size=10.5)



def build(data: dict, dst: Path) -> None:
    doc = Document()
    set_style(doc)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = Cm(1.6)
    sec.bottom_margin = Cm(1.5)
    sec.left_margin = Cm(1.7)
    sec.right_margin = Cm(1.5)

    hp = sec.header.paragraphs[0]
    hp.text = ""
    hr = hp.add_run(f"申论 · 批改讲义 · {data['title']}")
    hr.font.size = Pt(8.5)
    hr.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    hr.font.name = FONT
    hr._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    para(doc, data["title"], size=16, bold=True, space_after=2)
    para(doc, "考场版解析：题目 · 你的作答 · 得分 · 逐条对照 · 漏点 · 升格版 · 三步法", size=9, color=(0x66, 0x66, 0x66))

    emit_blocks(doc, data["block"])

    doc.save(str(dst))


def main() -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", default=str(SRC))
    ap.add_argument("--out", default=str(OUT))
    ap.add_argument("--name", default="")
    args = ap.parse_args()

    src = Path(args.src)
    data = parse(src)
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)

    m = re.search(r"(\d{4}-\d{2}-\d{2})", src.read_text(encoding="utf-8"))
    date = m.group(1) if m else "today"
    mm = re.search(r"练习\s*(\d+)", src.read_text(encoding="utf-8"))
    if mm:
        data["title"] = f"练习 {mm.group(1)} · 批改解析"
    name = args.name or f"{date}-review"

    docx = out / f"{name}.docx"
    build(data, docx)
    print(f"docx → {docx}")
    pdf = out / f"{name}.pdf"
    if to_pdf(docx, pdf):
        print(f"pdf  → {pdf}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
