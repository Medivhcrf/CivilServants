"""整卷摸底用的打印版：材料 + 若干道题，每题下面给足答题横线。

用法：
    python tools\\make_paper.py                      # 读 申论/今日任务.md
    python tools\\make_paper.py --src <md> --name 2026-10-31-paper
"""
import argparse
import math
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import make_handout as H  # noqa: E402
from docx.shared import Pt  # noqa: E402

ROOT = H.ROOT
OUT = H.OUT


def parse_paper(src: Path) -> dict:
    text = src.read_text(encoding="utf-8")
    lines = text.splitlines()

    date = re.search(r"(\d{4}-\d{2}-\d{2})", text[:600])
    day = re.search(r"Day\s*(\d+)", text[:600])
    week = re.search(r"（(周[一二三四五六日])）", text[:600])

    mats = []
    for m in re.finditer(r"\*\*材料([一二三四五六七八九十]+)\*\*\s*\n+(.*?)(?=\n\s*\*\*材料|\n---|\n## )",
                         text, re.S):
        body = H.clean(" ".join(x.strip() for x in m.group(2).splitlines() if x.strip()))
        if body:
            mats.append((f"材料{m.group(1)}", body))

    rows = []
    for ln in lines:
        if ln.strip().startswith("|") and ln.count("|") >= 3:
            cells = [H.clean(c) for c in ln.strip().strip("|").split("|")]
            if set("".join(cells)) <= set("-: "):
                continue
            rows.append(cells)
    sched = []
    if rows:
        start = next((i for i, r in enumerate(rows) if r and r[0] in ("时段", "时")), None)
        if start is not None:
            for r in rows[start + 1:]:
                if not r or not r[0] or r[0].startswith("轨道"):
                    break
                if not re.match(r"^\d{1,2}\s*[:：]\s*\d{2}", r[0]):
                    break
                sched.append(r)

    # 题目：### 第 N 题（15 分 · 250 字以内）  +  > **题目 N**：……
    questions = []
    meta = ""
    for ln in lines:
        s = ln.strip()
        m = re.match(r"^###\s*第\s*(\d+)\s*题[（(](.+?)[）)]", s)
        if m:
            meta = f"第 {m.group(1)} 题（{m.group(2)}）"
            continue
        m = re.match(r"^>\s*\*\*题目\s*(\d+)\*\*[:：]\s*(.+)$", s)
        if m:
            questions.append({"meta": meta or f"第 {m.group(1)} 题", "stem": H.clean(m.group(2))})
            meta = ""

    return {
        "date": date.group(1) if date else "0000-00-00",
        "day": day.group(1) if day else "?",
        "week": week.group(1) if week else "",
        "mats": mats,
        "sched": sched,
        "questions": questions,
    }


def answer_lines(meta: str) -> int:
    if "提纲" in meta:
        return 8
    m = re.search(r"(\d+)\s*字", meta)
    if not m:
        return 12
    n = int(m.group(1))
    return max(8, min(22, math.ceil(n / 26)))


def build(data: dict, dst: Path) -> None:
    header = f"申论 · Day {data['day']} · {data['date']}{data['week']} · 整卷摸底　　姓名／日期：__________"
    doc = H.new_doc(header)

    H.para(doc, f"Day {data['day']}　{data['date']}{data['week']}　整卷摸底", size=16, bold=True, space_after=2)
    H.para(doc, "不计时，但设硬截止；做不完就写“没做完”——没做完的位置就是摸底结果。",
           size=9, bold=True, color=(0xB0, 0x30, 0x30))

    if data["sched"]:
        H.heading(doc, "时间表")
        for r in data["sched"]:
            r = r + [""] * (4 - len(r))
            line = f"{r[0]}　{r[1]}　{r[2]}"
            if r[3]:
                line += f"　（{r[3]}）"
            H.para(doc, line, size=9, space_after=1.5, line=1.15)

    H.heading(doc, "一、材料（右边空白处写勾画）")
    for label, body in data["mats"]:
        t = doc.add_table(rows=1, cols=2)
        t.style = "Table Grid"
        t.autofit = False
        left, right = t.rows[0].cells
        H.cell_width(left, 14.0)
        H.cell_width(right, 3.4)
        left.text = ""
        p = left.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(label + "　")
        r.bold = True
        r.font.size = Pt(9)
        r.font.name = H.FONT
        r2 = p.add_run(body)
        r2.font.size = Pt(9)
        r2.font.name = H.FONT
        for row in t.rows:
            from docx.oxml import OxmlElement
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        sp = doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(0)
        sp.paragraph_format.line_spacing = 0.5

    H.heading(doc, "二、题目")
    for i, q in enumerate(data["questions"]):
        if i or data["mats"]:
            doc.add_page_break()
        H.heading(doc, q["meta"])
        H.para(doc, q["stem"], size=10, bold=True, space_after=6)
        H.para(doc, "答：", size=9.5, space_after=2)
        H.ruled(doc, answer_lines(q["meta"]))

    H.para(doc, "交卷：勾画稿 ＋ 各题答案 ＋ 一行“每题用时与卡点”。", size=8.5,
           color=(0x66, 0x66, 0x66), space_before=8)

    doc.save(str(dst))
    print(f"docx → {dst}")
    pdf = dst.with_suffix(".pdf")
    if H.to_pdf(dst, pdf):
        print(f"pdf  → {pdf}")
    else:
        print("pdf 转换失败", file=sys.stderr)


def main() -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", default=str(ROOT / "申论" / "今日任务.md"))
    ap.add_argument("--name", default="")
    args = ap.parse_args()
    data = parse_paper(Path(args.src))
    if not data["questions"]:
        print("没解析到题目（需要 ### 第 N 题（…） 与 > **题目 N**：…）", file=sys.stderr)
        return 1
    name = args.name or f"{data['date']}-paper"
    OUT.mkdir(parents=True, exist_ok=True)
    dst = OUT / f"{name}.docx"
    build(data, dst)
    import shutil
    shutil.copyfile(dst.with_suffix(".pdf"), OUT / "today.pdf")
    print(f"today.pdf ← {dst.with_suffix('.pdf').name}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
